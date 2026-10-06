#!/usr/bin/env python3
"""
Generate a dummy financial-services asset inventory seeded with software
versions that carry real, publicly known CVEs in the NIST NVD.

The workbook uses the same 'Asset Inventory' schema as
financial_services_dummy_asset_inventory.xlsx, so it can be fed straight into
nvd_asset_vulnerability_mvp.py. Vendor and 'Asset Subtype' values are chosen so
the agent's normalization (normalize_vendor / normalize_product) produces the
exact NVD CPE vendor:product pair.

An 'Expected Findings' sheet is the answer key: the CVE(s) each asset should be
flagged for, plus negative controls (patched versions and in-house software
that should NOT resolve to a CPE).

Usage:
    python scripts/generate_vulnerable_inventory.py \
        --output data/vulnerable_asset_inventory.xlsx
"""

from __future__ import annotations

import argparse
import functools
import importlib.util
import pathlib
import random
import sys
from datetime import date, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Asset ID", "Asset Name", "Asset Type", "Asset Subtype", "Asset Details",
    "Business Purpose", "Environment", "Deployment Location", "City / Region",
    "Technical Details", "Vendor", "Version / Firmware", "Operating System",
    "Hostname", "IP Address", "Network Zone", "Owner / Team",
    "Business Criticality", "Data Classification", "Asset Status",
    "Cloud / On-Prem", "High Availability", "Backup Required", "Monitoring",
    "Last Vulnerability Scan", "Next Lifecycle Review", "Primary Dependency",
]

LOCATIONS = {
    "Dallas DC1": "Dallas, TX",
    "Ashburn Cloud Region": "Ashburn, VA",
    "Charlotte DC1": "Charlotte, NC",
    "Charlotte DC2": "Charlotte, NC",
    "Phoenix DR": "Phoenix, AZ",
    "New York DC1": "New York, NY",
    "Chicago Colo": "Chicago, IL",
    "London Office": "London, UK",
}

PURPOSES = {
    "customer": "Supports customer-facing financial services workloads",
    "risk": "Supports regulatory, risk, and compliance processes",
    "infra": "Provides secure connectivity and infrastructure services",
    "data": "Provides data processing and analytics capabilities",
    "internal": "Supports internal operations and enterprise applications",
}

# (asset_type, subtype, vendor, version, os, role, purpose, env, zone, owner,
#  criticality, classification, dependency, expected_cves, note)
# expected_cves: real NVD CVEs the version falls inside. Empty = negative control.
ASSETS = [
    # --- Known-vulnerable web / middleware tier --------------------------------
    ("Server", "Apache HTTP Server", "Apache", "2.4.49", "Red Hat Enterprise Linux 8",
     "online-banking-web", "customer", "Production", "DMZ", "Application Engineering",
     "Critical", "Customer Financial Data", "API Gateway",
     ["CVE-2021-41773", "CVE-2021-42013"], "Path traversal / RCE; CISA KEV"),
    ("Server", "Apache HTTP Server", "Apache", "2.4.50", "Ubuntu 20.04 LTS",
     "mortgage-portal-web", "customer", "UAT", "DMZ", "Application Engineering",
     "High", "Sensitive Personal Data", "API Gateway",
     ["CVE-2021-42013"], "Incomplete fix for CVE-2021-41773; CISA KEV"),
    ("Application", "Log4j", "Apache", "2.14.1", "Red Hat Enterprise Linux 8",
     "payments-ledger-svc", "customer", "Production", "PCI", "Payments Technology",
     "Critical", "PCI Cardholder Data", "Message Bus",
     ["CVE-2021-44228", "CVE-2021-45046"], "Log4Shell; CISA KEV"),
    ("Server", "Tomcat", "Apache", "9.0.30", "Red Hat Enterprise Linux 7",
     "loan-origination-app", "customer", "Production", "Server", "Application Engineering",
     "High", "Sensitive Personal Data", "Database Cluster",
     ["CVE-2020-1938"], "Ghostcat AJP file read/inclusion; CISA KEV"),
    ("Application", "Struts", "Apache", "2.3.31", "Red Hat Enterprise Linux 7",
     "legacy-claims-portal", "risk", "Production", "DMZ", "Risk Technology",
     "High", "Confidential", "Database Cluster",
     ["CVE-2017-5638"], "Content-Type OGNL RCE (Equifax); CISA KEV"),
    ("Application", "ActiveMQ", "Apache", "5.18.2", "Ubuntu 22.04 LTS",
     "trade-events-broker", "data", "Production", "Server", "DevOps Platform",
     "Critical", "Customer Financial Data", "Message Bus",
     ["CVE-2023-46604"], "OpenWire deserialization RCE; CISA KEV"),
    ("Application", "OFBiz", "Apache", "18.12.10", "Ubuntu 22.04 LTS",
     "procurement-erp", "internal", "Production", "Corporate", "Corporate IT",
     "Medium", "Internal", "Database Cluster",
     ["CVE-2023-51467"], "Authentication bypass leading to SSRF/RCE; CISA KEV"),
    ("Application", "Spring Framework", "VMware", "5.3.17", "Ubuntu 20.04 LTS",
     "card-rewards-api", "customer", "Production", "PCI", "Payments Technology",
     "Critical", "PCI Cardholder Data", "API Gateway",
     ["CVE-2022-22965"], "Spring4Shell data-binding RCE; CISA KEV"),
    ("Server", "WebLogic Server", "Oracle", "12.2.1.3.0", "Oracle Linux 7",
     "treasury-settlement-app", "customer", "Production", "Server", "Application Engineering",
     "Critical", "Customer Financial Data", "Database Cluster",
     ["CVE-2020-14882"], "Console unauthenticated RCE; CISA KEV"),
    ("Server", "nginx", "F5", "1.20.0", "Ubuntu 20.04 LTS",
     "mobile-api-edge", "customer", "Production", "DMZ", "Cloud Engineering",
     "High", "Customer Financial Data", "API Gateway",
     ["CVE-2021-23017"], "DNS resolver off-by-one"),
    # --- Collaboration / DevOps / file transfer --------------------------------
    ("Application", "Confluence Server", "Atlassian", "7.18.0", "Ubuntu 20.04 LTS",
     "eng-wiki", "internal", "Production", "Corporate", "Corporate IT",
     "Medium", "Confidential", "Identity Platform",
     ["CVE-2022-26134"], "OGNL injection RCE; CISA KEV"),
    ("Application", "MOVEit Transfer", "Progress", "2023.0.0", "Windows Server 2019",
     "partner-file-exchange", "customer", "Production", "DMZ", "Core Infrastructure",
     "Critical", "Sensitive Personal Data", "Storage Platform",
     ["CVE-2023-34362"], "SQL injection exploited by Cl0p; CISA KEV"),
    ("Application", "Jenkins", "Jenkins", "2.441", "Ubuntu 22.04 LTS",
     "ci-build-controller", "internal", "Production", "Management", "DevOps Platform",
     "High", "Confidential", "Identity Platform",
     ["CVE-2024-23897"], "CLI arbitrary file read; CISA KEV"),
    ("Application", "TeamCity", "JetBrains", "2023.05.3", "Ubuntu 22.04 LTS",
     "release-pipeline", "internal", "Production", "Management", "DevOps Platform",
     "High", "Confidential", "Identity Platform",
     ["CVE-2023-42793"], "Authentication bypass RCE; CISA KEV"),
    ("Application", "GitLab", "GitLab", "16.1.0", "Ubuntu 22.04 LTS",
     "source-control", "internal", "Production", "Management", "DevOps Platform",
     "High", "Restricted", "Identity Platform",
     ["CVE-2023-7028"], "Account takeover via password reset; CISA KEV"),
    ("Application", "Grafana", "Grafana", "8.3.0", "Ubuntu 20.04 LTS",
     "noc-dashboards", "infra", "Production", "Management", "Network Engineering",
     "Medium", "Internal", "Network Core",
     ["CVE-2021-43798"], "Plugin path traversal; CISA KEV"),
    ("Application", "Drupal", "Drupal", "7.57", "CentOS 7",
     "marketing-microsite", "customer", "Production", "DMZ", "Application Engineering",
     "Low", "Public", "API Gateway",
     ["CVE-2018-7600"], "Drupalgeddon2 RCE; CISA KEV"),
    ("Server", "Exim", "Exim", "4.91", "Debian 9",
     "smtp-relay", "infra", "Production", "DMZ", "Core Infrastructure",
     "Medium", "Internal", "DNS/DHCP",
     ["CVE-2019-10149"], "Return of the WIZard RCE; CISA KEV"),
    # --- Crypto / remote access libraries --------------------------------------
    ("Application", "OpenSSL", "OpenSSL", "3.0.6", "Ubuntu 22.04 LTS",
     "hsm-gateway", "customer", "Production", "PCI", "Cybersecurity",
     "Critical", "PCI Cardholder Data", "Identity Platform",
     ["CVE-2022-3602", "CVE-2022-3786"], "X.509 punycode buffer overflows"),
    ("Server", "OpenSSH", "OpenBSD", "8.9p1", "Ubuntu 22.04 LTS",
     "jump-host", "infra", "Production", "Management", "Cybersecurity",
     "High", "Restricted", "Identity Platform",
     ["CVE-2024-6387"], "regreSSHion signal-handler race RCE"),
    # --- Data tier -------------------------------------------------------------
    ("Database", "Redis Cache", "Redis Labs", "7.0.10", "Ubuntu 22.04 LTS",
     "session-cache", "customer", "Production", "Server", "Cloud Engineering",
     "High", "Sensitive Personal Data", "Database Cluster",
     ["CVE-2023-28856"], "HINCRBYFLOAT crash (DoS)"),
    ("Database", "PostgreSQL Database", "PostgreSQL", "13.2", "Red Hat Enterprise Linux 8",
     "kyc-records-db", "risk", "Production", "Restricted", "Data Engineering",
     "Critical", "Sensitive Personal Data", "Storage Platform",
     ["CVE-2021-32027"], "Array subscript buffer overrun"),
    ("Database", "Elasticsearch", "Elastic", "7.17.12", "Ubuntu 20.04 LTS",
     "fraud-search-index", "data", "Production", "Restricted", "Data Engineering",
     "High", "Customer Financial Data", "Storage Platform",
     ["CVE-2023-31419"], "_search API stack overflow (DoS)"),
    # --- DR / non-prod copies of vulnerable software ---------------------------
    ("Application", "Log4j", "Apache", "2.14.1", "Red Hat Enterprise Linux 8",
     "payments-ledger-svc", "customer", "DR", "PCI", "Payments Technology",
     "Critical", "PCI Cardholder Data", "Message Bus",
     ["CVE-2021-44228", "CVE-2021-45046"], "DR replica of production ledger"),
    ("Application", "Confluence Server", "Atlassian", "7.18.0", "Ubuntu 20.04 LTS",
     "eng-wiki", "internal", "Sandbox", "Corporate", "Corporate IT",
     "Low", "Internal", "Identity Platform",
     ["CVE-2022-26134"], "Sandbox copy - same CVE, lower business risk"),
    # --- Negative controls: patched versions (should NOT flag listed CVE) -----
    ("Application", "Log4j", "Apache", "2.17.1", "Red Hat Enterprise Linux 9",
     "statements-svc", "customer", "Production", "Server", "Application Engineering",
     "High", "Customer Financial Data", "Message Bus",
     [], "CONTROL: patched; must not flag CVE-2021-44228/45046"),
    ("Server", "Tomcat", "Apache", "9.0.31", "Red Hat Enterprise Linux 9",
     "branch-locator-app", "customer", "Production", "DMZ", "Application Engineering",
     "Medium", "Public", "API Gateway",
     [], "CONTROL: first fixed release; must not flag CVE-2020-1938"),
    ("Application", "Jenkins", "Jenkins", "2.442", "Ubuntu 22.04 LTS",
     "ci-build-agent-pool", "internal", "Development", "Management", "DevOps Platform",
     "Low", "Internal", "Identity Platform",
     [], "CONTROL: fixed release; must not flag CVE-2024-23897"),
    # --- Negative controls: in-house software with no NVD CPE -----------------
    ("Application", "Core Banking Ledger", "Internal", "5.2.0", "Red Hat Enterprise Linux 9",
     "core-banking", "customer", "Production", "Restricted", "Application Engineering",
     "Critical", "Customer Financial Data", "Database Cluster",
     [], "CONTROL: proprietary; expect 'No strong CPE candidate'"),
    ("Application", "Branch Teller UI", "Internal", "3.4", "Windows 11 Enterprise",
     "teller-ui", "customer", "Production", "Corporate", "Application Engineering",
     "High", "Customer Financial Data", "API Gateway",
     [], "CONTROL: proprietary; expect 'No strong CPE candidate'"),
]


def tech_details(asset_type: str, subtype: str, version: str, os_name: str, rng: random.Random) -> str:
    cpu = rng.choice([8, 16, 24, 32, 48])
    ram = rng.choice([16, 32, 64, 128])
    if asset_type == "Database":
        return (f"Engine: {subtype} {version}; Nodes: {rng.choice([1, 3, 5])}; vCPU: {cpu}; "
                f"RAM: {ram} GB; Storage: {rng.choice([500, 1000, 4000])} GB; "
                "Encryption: AES-256 at rest; TLS: 1.2/1.3")
    if asset_type == "Server":
        return (f"Software: {subtype} {version}; CPU: {cpu} vCPU; RAM: {ram} GB; OS: {os_name}; "
                f"Hypervisor: {rng.choice(['VMware ESXi 8.0', 'Hyper-V 2022', 'KVM'])}; NIC: 10GbE")
    return (f"Software: {subtype}; Version: {version}; Runtime: "
            f"{rng.choice(['Java 11', 'Java 17', 'Python 3.11', 'Node 18', '.NET 6'])}; "
            f"OS: {os_name}; Deployment: {rng.choice(['VMware', 'Kubernetes', 'Bare metal'])}")


def build(seed: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    rng = random.Random(seed)
    today = date(2026, 10, 6)
    rows, key = [], []
    env_code = {"Production": "prd", "DR": "dr", "UAT": "uat", "Development": "dev",
                "Sandbox": "sbx", "Test": "tst"}

    for i, (atype, subtype, vendor, version, os_name, role, purpose, env, zone, owner,
            crit, cls, dep, cves, note) in enumerate(ASSETS, start=1):
        asset_id = f"VA-{i:05d}"
        name = f"{role}-{env_code[env]}-{i:04d}"
        location = rng.choice(list(LOCATIONS))
        if env == "DR":
            location = "Phoenix DR"
        rows.append({
            "Asset ID": asset_id,
            "Asset Name": name,
            "Asset Type": atype,
            "Asset Subtype": subtype,
            "Asset Details": (f"{subtype} {version} supporting {role.replace('-', ' ')}. "
                              f"Primary owner is {owner}. Located in {location}."),
            "Business Purpose": PURPOSES[purpose],
            "Environment": env,
            "Deployment Location": location,
            "City / Region": LOCATIONS[location],
            "Technical Details": tech_details(atype, subtype, version, os_name, rng),
            "Vendor": vendor,
            "Version / Firmware": version,
            "Operating System": os_name,
            "Hostname": name,
            "IP Address": f"10.{20 + (i // 250)}.{rng.randint(0, 254)}.{i % 250 + 1}",
            "Network Zone": zone,
            "Owner / Team": owner,
            "Business Criticality": crit,
            "Data Classification": cls,
            "Asset Status": "Active",
            "Cloud / On-Prem": rng.choice(["Cloud", "Hybrid", "On-Premises"]),
            "High Availability": "Yes" if crit in ("Critical", "High") else "No",
            "Backup Required": "Yes",
            "Monitoring": "24x7" if crit == "Critical" else rng.choice(["Automated Alerts", "Business Hours"]),
            "Last Vulnerability Scan": today - timedelta(days=rng.randint(30, 400)),
            "Next Lifecycle Review": today + timedelta(days=rng.randint(90, 900)),
            "Primary Dependency": dep,
        })
        key.append({
            "Asset ID": asset_id,
            "Asset Name": name,
            "Vendor": vendor,
            "Product": subtype,
            "Version": version,
            "Expected NVD CPE (vendor:product)": _expected_cpe(vendor, subtype),
            "Expected CVEs": ", ".join(cves) if cves else "None (control)",
            "Should Be Flagged": "Yes" if cves else "No",
            "Notes": note,
        })
    return pd.DataFrame(rows, columns=COLUMNS), pd.DataFrame(key)


@functools.lru_cache(maxsize=1)
def _agent_module():
    # Reuse the agent's own normalization so the answer key shows what it will query.
    path = pathlib.Path(__file__).resolve().parent.parent / "nvd_asset_vulnerability_mvp.py"
    spec = importlib.util.spec_from_file_location("nvd_mvp", path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)  # type: ignore[union-attr]
    return mod


def _expected_cpe(vendor: str, subtype: str) -> str:
    mod = _agent_module()
    return f"{mod.normalize_vendor(vendor)}:{mod.normalize_product(subtype)}"


def format_workbook(path: str) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F3864")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for idx, col in enumerate(ws.columns, start=1):
            width = max(len("" if c.value is None else str(c.value)) for c in col)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 60)
        # Keep versions as text so '7.0' / '2.4.50' are not coerced to numbers.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.number_format = "@"
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="data/vulnerable_asset_inventory.xlsx")
    parser.add_argument("--seed", type=int, default=2026)
    args = parser.parse_args()

    inventory, key = build(args.seed)
    dictionary = pd.DataFrame([
        ["Asset ID", "Unique synthetic identifier (VA- prefix = vulnerability test set)"],
        ["Asset Subtype", "Product name; chosen to normalize to the NVD CPE product"],
        ["Vendor", "Vendor; chosen to normalize to the NVD CPE vendor"],
        ["Version / Firmware", "Installed version; intentionally set inside known-vulnerable ranges"],
        ["Expected Findings", "Answer key sheet (not read by the agent)"],
    ], columns=["Field", "Description"])

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        inventory.to_excel(writer, sheet_name="Asset Inventory", index=False)
        key.to_excel(writer, sheet_name="Expected Findings", index=False)
        dictionary.to_excel(writer, sheet_name="Data Dictionary", index=False)
    format_workbook(args.output)

    print(f"Assets: {len(inventory)}  "
          f"(expected vulnerable: {(key['Should Be Flagged'] == 'Yes').sum()}, "
          f"controls: {(key['Should Be Flagged'] == 'No').sum()})")
    print(f"Output: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
